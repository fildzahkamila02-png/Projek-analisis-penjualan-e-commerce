"""Ringkasan Excel dengan rumus (Materi 4): data bersih + analisis SUMIFS/COUNTIF/XLOOKUP-style."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

df = pd.read_csv("data/clean/transaksi_clean.csv")
produk = pd.read_csv("data/raw/produk.csv")
df = df.merge(produk[["ID_Produk", "Kategori"]], on="ID_Produk")
df = df[["ID_Transaksi", "Tanggal", "Bulan", "ID_Produk", "Kategori", "Kota", "Qty", "Harga", "Total", "Metode_Bayar"]]

wb = Workbook()
HEAD = PatternFill("solid", fgColor="2A6FBB")
def header(ws, row, kolom):
    for j, k in enumerate(kolom, 1):
        c = ws.cell(row=row, column=j, value=k)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD

# Sheet Data
ws = wb.active
ws.title = "Data"
header(ws, 1, list(df.columns))
for r in df.itertuples(index=False):
    ws.append(list(r))
n = len(df) + 1
for j in range(1, 11):
    ws.column_dimensions[get_column_letter(j)].width = 15
ws.freeze_panes = "A2"

# Sheet Ringkasan (rumus ke sheet Data)
rg = wb.create_sheet("Ringkasan")
rg["A1"], rg["A1"].font = "Ringkasan Penjualan 2025 (rumus Excel)", Font(bold=True, size=14)
for i, (label, rumus) in enumerate([
        ("Jumlah transaksi", f"=COUNTA(Data!A2:A{n})"),
        ("Total omzet (Rp)", f"=SUM(Data!I2:I{n})"),
        ("Rata-rata per transaksi", "=B4/B3"),
        ("Median per transaksi", f"=MEDIAN(Data!I2:I{n})"),
        ("Transaksi grosir (Qty>=50)", f'=COUNTIF(Data!G2:G{n},">=50")')], 3):
    rg.cell(row=i, column=1, value=label)
    rg.cell(row=i, column=2, value=rumus)

def tabel(ws, baris, judul, kolom_data, daftar, kolom_nilai="I"):
    ws.cell(row=baris, column=1, value=judul).font = Font(bold=True)
    header(ws, baris + 1, [kolom_data, "Omzet (Rp)", "Jumlah Transaksi", "% Omzet"])
    for k, nama in enumerate(daftar, baris + 2):
        ws.cell(row=k, column=1, value=nama)
        ws.cell(row=k, column=2, value=f"=SUMIFS(Data!${kolom_nilai}$2:${kolom_nilai}${n},Data!${KOL[kolom_data]}$2:${KOL[kolom_data]}${n},A{k})")
        ws.cell(row=k, column=3, value=f"=COUNTIFS(Data!${KOL[kolom_data]}$2:${KOL[kolom_data]}${n},A{k})")
        ws.cell(row=k, column=4, value=f"=B{k}/$B$4").number_format = "0.0%"
    return baris + 3 + len(daftar)

KOL = {"Bulan": "C", "Kategori": "E", "Kota": "F", "Metode_Bayar": "J"}
b = 10
for kolom in ["Bulan", "Kategori", "Kota", "Metode_Bayar"]:
    b = tabel(rg, b, f"Omzet per {kolom}", kolom, sorted(df[kolom].unique()))

# Sheet Lookup (XLOOKUP/VLOOKUP-style): harga dasar produk
lk = wb.create_sheet("Produk")
header(lk, 1, list(produk.columns))
for r in produk.itertuples(index=False):
    lk.append(list(r))
lk["F1"], lk["G1"] = "Cari ID produk:", "P03"
lk["F2"], lk["G2"] = "Nama produk:", '=IFERROR(VLOOKUP(G1,A2:D17,2,FALSE),"Tidak ditemukan")'
lk["F3"], lk["G3"] = "Harga dasar:", '=IFERROR(VLOOKUP(G1,A2:D17,4,FALSE),"-")'
for w in (rg, lk):
    for j in range(1, 8):
        w.column_dimensions[get_column_letter(j)].width = 24
for row in rg.iter_rows(min_row=3, min_col=2, max_col=3):
    for c in row:
        c.number_format = "#,##0"
wb.save("output/xlsx/ringkasan_penjualan.xlsx")
print("OK xlsx")

# CSV siap upload ke Looker Studio (Materi 7)
looker = df.copy()
looker["Bulan_Angka"] = looker["Tanggal"].str[5:7].astype(int)
looker.to_csv("output/data_untuk_looker_studio.csv", index=False)
